import base64

from conftest import FakeChatClient, make_pdf

from src.core import vision
from src.core.vision import Page, VisionProvider, read_document, read_pages, resolve_vision_provider


# ---------- provider resolution ----------
def test_no_provider_when_no_keys():
    assert resolve_vision_provider({}) is None


def test_vision_api_key_wins_over_everything():
    p = resolve_vision_provider({
        "VISION_API_KEY": "v", "VISION_BASE_URL": "https://example.test/v1", "VISION_MODEL": "my-vl",
        "OPENROUTER_API_KEY": "r", "OPENAI_API_KEY": "o",
    })
    assert (p.name, p.api_key, p.base_url, p.model) == ("custom", "v", "https://example.test/v1", "my-vl")


def test_openrouter_before_openai_with_defaults():
    p = resolve_vision_provider({"OPENROUTER_API_KEY": "r", "OPENAI_API_KEY": "o"})
    assert p.name == "openrouter"
    assert p.base_url == "https://openrouter.ai/api/v1"
    assert p.model == "qwen/qwen2.5-vl-72b-instruct"


def test_openai_fallback_defaults():
    p = resolve_vision_provider({"OPENAI_API_KEY": "o"})
    assert (p.name, p.base_url, p.model) == ("openai", None, "gpt-4o-mini")


def test_vision_model_overrides_provider_default():
    p = resolve_vision_provider({"OPENROUTER_API_KEY": "r", "VISION_MODEL": "qwen/other"})
    assert p.model == "qwen/other"


def test_blank_keys_are_ignored():
    assert resolve_vision_provider({"VISION_API_KEY": "  ", "OPENAI_API_KEY": ""}) is None


# ---------- text-layer fallback ----------
def test_text_layer_fallback_without_key(tmp_path, monkeypatch, caplog):
    pdf = make_pdf(tmp_path / "deck.pdf", ["Binary search halves the range", "Hash tables use buckets"])

    def boom(*a, **k):
        raise AssertionError("no vision client may be created without a key")

    monkeypatch.setattr(vision, "make_vision_client", boom)
    with caplog.at_level("WARNING"):
        pages = read_document(pdf, tmp_path / "work", dpi=72)
    assert "Binary search halves the range" in pages[0]
    assert "Hash tables use buckets" in pages[1]
    assert "No vision provider configured" in caplog.text
    assert not (tmp_path / "work" / "pages").exists()  # nothing rendered in degraded mode


def test_pptx_text_layer_includes_notes(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Recursion"
    slide.placeholders[1].text = "Base case stops the recursion"
    slide.notes_slide.notes_text_frame.text = "Mention stack overflow"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    pages = read_document(path, tmp_path / "work", dpi=72)
    assert len(pages) == 1
    assert "Recursion" in pages[0] and "Base case stops the recursion" in pages[0]
    assert "Mention stack overflow" in pages[0]


# ---------- vision path (mocked client) ----------
PROVIDER = VisionProvider("openai", "test-key", None, "gpt-4o-mini")


def test_vision_path_sends_png_and_returns_transcription(tmp_path):
    pdf = make_pdf(tmp_path / "deck.pdf", ["Slide one", "Slide two"])
    client = FakeChatClient(lambda kw: "# Transcribed\n```python\nprint(1)\n```\n[Diagram: a tree]")

    pages = read_document(pdf, tmp_path / "work", dpi=72, provider=PROVIDER, client=client)

    assert pages == ["# Transcribed\n```python\nprint(1)\n```\n[Diagram: a tree]"] * 2
    assert len(client.calls) == 2
    call = client.calls[0]
    assert call["model"] == "gpt-4o-mini"
    user_parts = call["messages"][1]["content"]
    image_urls = [p["image_url"]["url"] for p in user_parts if p["type"] == "image_url"]
    assert len(image_urls) == 1 and image_urls[0].startswith("data:image/png;base64,")
    assert base64.b64decode(image_urls[0].split(",", 1)[1])[:8] == b"\x89PNG\r\n\x1a\n"
    assert "Slide one" in user_parts[0]["text"] or "Slide two" in user_parts[0]["text"]


def test_failed_page_retries_then_falls_back_to_text_layer(monkeypatch):
    monkeypatch.setenv("VISION_MAX_RETRIES", "2")
    attempts = {"n": 0}

    def flaky(kw):
        text = kw["messages"][1]["content"][0]["text"]
        if text.startswith("Page 2"):
            attempts["n"] += 1
            raise TimeoutError("simulated timeout")
        return "vision text"

    pages = [Page(1, "layer one", [b"png"]), Page(2, "layer two", [b"png"]), Page(3, "layer three", [])]
    out = read_pages(pages, PROVIDER, client=FakeChatClient(flaky))
    assert out == ["vision text", "layer two", "layer three"]
    assert attempts["n"] == 3  # 1 try + 2 retries


def test_page_cap_limits_vision_calls(tmp_path, monkeypatch):
    monkeypatch.setenv("VISION_MAX_PAGES", "2")
    pdf = make_pdf(tmp_path / "deck.pdf", ["p1", "p2", "p3"])
    client = FakeChatClient(lambda kw: "vision")
    pages = read_document(pdf, tmp_path / "work", dpi=72, provider=PROVIDER, client=client)
    assert pages[:2] == ["vision", "vision"]
    assert "p3" in pages[2]
    assert len(client.calls) == 2


def test_concurrency_is_bounded(monkeypatch):
    import threading
    import time

    monkeypatch.setenv("VISION_CONCURRENCY", "2")
    live = {"now": 0, "peak": 0}
    lock = threading.Lock()

    def slow(kw):
        with lock:
            live["now"] += 1
            live["peak"] = max(live["peak"], live["now"])
        time.sleep(0.05)
        with lock:
            live["now"] -= 1
        return "ok"

    pages = [Page(i, "", [b"png"]) for i in range(1, 7)]
    assert read_pages(pages, PROVIDER, client=FakeChatClient(slow)) == ["ok"] * 6
    assert live["peak"] <= 2
