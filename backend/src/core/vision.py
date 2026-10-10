# -*- coding: utf-8 -*-
"""
Vision-LLM page reader for lecture slides (replaces EasyOCR).

- Renders PDF pages (PyMuPDF) and PPTX slides (LibreOffice/COM when available,
  otherwise the slide's embedded pictures) to PNG.
- Asks an OpenAI-compatible vision model to transcribe each page: slide text,
  code blocks, and a short description of diagrams/charts.
- Falls back to the document's embedded text layer (PyMuPDF get_text /
  python-pptx) when no vision provider is configured, for pages past the page
  cap, and for any page whose vision call fails.

Provider resolution (first match wins):
  1) VISION_API_KEY  (+ optional VISION_BASE_URL, VISION_MODEL)
  2) OPENROUTER_API_KEY -> https://openrouter.ai/api/v1, qwen/qwen2.5-vl-72b-instruct
  3) OPENAI_API_KEY     -> OpenAI, gpt-4o-mini
  4) none               -> text layer only (degraded, logged as a warning)
"""

from __future__ import annotations

import base64
import io
import logging
import os
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, List, Mapping, Optional

log = logging.getLogger(__name__)

OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
OPENROUTER_DEFAULT_MODEL = "qwen/qwen2.5-vl-72b-instruct"
OPENAI_DEFAULT_MODEL = "gpt-4o-mini"

VISION_SYSTEM_PROMPT = (
    "You transcribe one lecture slide (or document page) for a quiz generator.\n"
    "- Transcribe ALL visible text faithfully in reading order as Markdown: the slide title as a "
    "'# ' heading, bullets as '- ' lists (keep nesting), numbered lists as numbered lists.\n"
    "- Reproduce code exactly (indentation, symbols, comments) inside fenced code blocks with a "
    "language tag when obvious.\n"
    "- Reproduce tables as Markdown tables. Write formulas as plain text or LaTeX.\n"
    "- For each diagram, chart, graph or figure add one short line '[Diagram: ...]' describing what it "
    "shows (labelled parts, axes, trends, relationships, key values).\n"
    "- Do not summarise, explain, translate, or add facts that are not on the page. Skip logos, page "
    "numbers and decorative elements.\n"
    "- The embedded text layer is given as a hint and may be incomplete or out of order; trust the image.\n"
    "- Return only the transcription. If the page has no content, return an empty string."
)


# =========================
# Settings
# =========================
def _env_int(name: str, default: int) -> int:
    try:
        return int(os.getenv(name, str(default)))
    except ValueError:
        return default


def _env_float(name: str, default: float) -> float:
    try:
        return float(os.getenv(name, str(default)))
    except ValueError:
        return default


def vision_settings() -> dict:
    """Read tunables at call time so tests and deployments can override via env."""
    return {
        "concurrency": max(1, _env_int("VISION_CONCURRENCY", 4)),
        "timeout": max(5.0, _env_float("VISION_PAGE_TIMEOUT_SECONDS", 60.0)),
        "retries": max(0, _env_int("VISION_MAX_RETRIES", 2)),
        "backoff": max(0.0, _env_float("VISION_RETRY_BACKOFF_SECONDS", 1.5)),
        "max_pages": max(1, _env_int("VISION_MAX_PAGES", 40)),
        "max_images_per_page": max(1, _env_int("VISION_MAX_IMAGES_PER_PAGE", 4)),
    }


# =========================
# Provider resolution
# =========================
@dataclass(frozen=True)
class VisionProvider:
    name: str
    api_key: str
    base_url: Optional[str]
    model: str


def resolve_vision_provider(env: Optional[Mapping[str, str]] = None) -> Optional[VisionProvider]:
    env = os.environ if env is None else env
    model_override = (env.get("VISION_MODEL") or "").strip()

    key = (env.get("VISION_API_KEY") or "").strip()
    if key:
        base_url = (env.get("VISION_BASE_URL") or "").strip() or None
        return VisionProvider("custom", key, base_url, model_override or OPENAI_DEFAULT_MODEL)

    key = (env.get("OPENROUTER_API_KEY") or "").strip()
    if key:
        return VisionProvider("openrouter", key, OPENROUTER_BASE_URL, model_override or OPENROUTER_DEFAULT_MODEL)

    key = (env.get("OPENAI_API_KEY") or "").strip()
    if key:
        return VisionProvider("openai", key, None, model_override or OPENAI_DEFAULT_MODEL)

    return None


def make_vision_client(provider: VisionProvider, timeout: float):
    """OpenAI-compatible client. Retries are handled by read_pages, not the SDK."""
    from openai import OpenAI

    kwargs: dict = {"api_key": provider.api_key, "timeout": timeout, "max_retries": 0}
    if provider.base_url:
        kwargs["base_url"] = provider.base_url
    return OpenAI(**kwargs)


# =========================
# Page model
# =========================
@dataclass
class Page:
    index: int                                          # 1-based
    text: str = ""                                      # embedded text layer
    images: List[bytes] = field(default_factory=list)   # PNG bytes sent to the vision model
    notes: str = ""                                     # speaker notes (PPTX), appended verbatim


def _with_notes(page: Page, body: str) -> str:
    body = (body or "").strip()
    notes = (page.notes or "").strip()
    if notes:
        body = f"{body}\n\n> Notes:\n> {notes}".strip()
    return body


def _to_png(blob: bytes, max_side: int = 2000, min_side: int = 64) -> Optional[bytes]:
    """Normalise an embedded picture to PNG; skip icons/bullets that are too small to matter."""
    try:
        from PIL import Image

        img = Image.open(io.BytesIO(blob))
        if min(img.size) < min_side:
            return None
        if img.mode not in ("RGB", "RGBA", "L"):
            img = img.convert("RGBA")
        img.thumbnail((max_side, max_side))
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()
    except Exception:
        return None


# =========================
# Document loaders
# =========================
def load_pdf_pages(path: str | Path, work_dir: str | Path, dpi: int, render: bool) -> List[Page]:
    import fitz  # PyMuPDF
    from .render import _pdf_to_png_pymupdf

    with fitz.open(str(path)) as doc:
        texts = [page.get_text().strip() for page in doc]

    pngs: List[Path] = []
    if render and texts:
        max_pages = vision_settings()["max_pages"]
        pngs = _pdf_to_png_pymupdf(path, Path(work_dir) / "pages", dpi=dpi, max_pages=max_pages)

    pages: List[Page] = []
    for i, text in enumerate(texts, 1):
        images = [pngs[i - 1].read_bytes()] if i <= len(pngs) else []
        pages.append(Page(index=i, text=text, images=images))
    return pages


def _pptx_slide_text_and_pictures(slide) -> tuple[str, List[bytes], str]:
    chunks: List[str] = []
    pictures: List[bytes] = []
    for shape in slide.shapes:
        try:
            if getattr(shape, "has_text_frame", False) or hasattr(shape, "text"):
                if shape.text:
                    chunks.append(shape.text)
        except Exception:
            pass
        try:
            image = getattr(shape, "image", None)
            if image is not None:
                pictures.append(image.blob)
        except Exception:
            pass
    notes = ""
    try:
        if slide.has_notes_slide and slide.notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text or ""
    except Exception:
        pass
    return "\n\n".join(chunks), pictures, notes


def load_pptx_pages(path: str | Path, work_dir: str | Path, dpi: int, render: bool) -> List[Page]:
    from pptx import Presentation

    settings = vision_settings()
    prs = Presentation(str(path))

    slide_pngs: List[Path] = []
    if render:
        # Full-slide render needs LibreOffice (or PowerPoint COM on Windows); the
        # Docker image has neither, so this usually falls through to embedded pictures.
        try:
            from .render import render_slides_to_images

            slide_pngs = sorted(render_slides_to_images(
                str(path), out_dir=str(Path(work_dir) / "slides"), dpi=dpi, prefer_windows_com=True
            ))
        except Exception as e:
            log.info("PPTX slide rendering unavailable; sending embedded pictures instead (%s)", str(e).splitlines()[0])
            slide_pngs = []

    pages: List[Page] = []
    for i, slide in enumerate(prs.slides, 1):
        text, pictures, notes = _pptx_slide_text_and_pictures(slide)
        images: List[bytes] = []
        if render and i <= settings["max_pages"]:
            if len(slide_pngs) >= i:
                images = [slide_pngs[i - 1].read_bytes()]
            else:
                for blob in pictures:
                    png = _to_png(blob)
                    if png:
                        images.append(png)
                    if len(images) >= settings["max_images_per_page"]:
                        break
        pages.append(Page(index=i, text=text, images=images, notes=notes))
    return pages


# =========================
# Vision calls
# =========================
def _data_url(png: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(png).decode("ascii")


def _transcribe_page(client: Any, provider: VisionProvider, page: Page, language: str, timeout: float) -> str:
    hint = (page.text or "").strip()[:4000] or "(empty)"
    content: List[dict] = [{
        "type": "text",
        "text": (
            f"Page {page.index}. Document language: {language or 'en'}.\n"
            f"Embedded text layer (hint):\n{hint}"
        ),
    }]
    content.extend({"type": "image_url", "image_url": {"url": _data_url(png)}} for png in page.images)
    resp = client.chat.completions.create(
        model=provider.model,
        messages=[
            {"role": "system", "content": VISION_SYSTEM_PROMPT},
            {"role": "user", "content": content},
        ],
        temperature=0,
        timeout=timeout,
    )
    return (resp.choices[0].message.content or "").strip() if resp.choices else ""


def _read_one(client: Any, provider: VisionProvider, page: Page, language: str, settings: dict) -> str:
    if not page.images:
        return _with_notes(page, page.text)

    last_error: Optional[Exception] = None
    for attempt in range(settings["retries"] + 1):
        try:
            text = _transcribe_page(client, provider, page, language, settings["timeout"])
            if text:
                return _with_notes(page, text)
            last_error = RuntimeError("empty response")
        except Exception as e:  # network, timeout, rate limit, bad response
            last_error = e
        if attempt < settings["retries"]:
            time.sleep(settings["backoff"] * (2 ** attempt))

    log.warning(
        "Vision read failed for page %d after %d attempt(s) (%s); using the text layer for this page.",
        page.index, settings["retries"] + 1, type(last_error).__name__ if last_error else "unknown",
    )
    return _with_notes(page, page.text)


def read_pages(
    pages: List[Page],
    provider: Optional[VisionProvider],
    language: str = "en",
    client: Any = None,
) -> List[str]:
    """Return one transcription per page, in page order."""
    if provider is None:
        return [_with_notes(p, p.text) for p in pages]

    settings = vision_settings()
    if client is None:
        client = make_vision_client(provider, settings["timeout"])
    with ThreadPoolExecutor(max_workers=settings["concurrency"]) as pool:
        return list(pool.map(lambda p: _read_one(client, provider, p, language, settings), pages))


def read_document(
    path: str | Path,
    work_dir: str | Path,
    dpi: int = 180,
    language: str = "en",
    provider: Optional[VisionProvider] = None,
    client: Any = None,
) -> List[str]:
    """
    Read a .pdf or .pptx into per-page Markdown.
    `provider` defaults to resolve_vision_provider(); pass a client to inject a fake in tests.
    """
    if provider is None:
        provider = resolve_vision_provider()
    if provider is None:
        log.warning(
            "No vision provider configured (set VISION_API_KEY, OPENROUTER_API_KEY or OPENAI_API_KEY). "
            "Falling back to the document's embedded text layer; scanned pages and diagrams will be missed."
        )

    ext = Path(path).suffix.lower()
    dpi = max(72, min(int(dpi), 300))
    render = provider is not None
    if ext == ".pdf":
        pages = load_pdf_pages(path, work_dir, dpi, render)
    elif ext == ".pptx":
        pages = load_pptx_pages(path, work_dir, dpi, render)
    else:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}. Upload .pptx or .pdf")

    return read_pages(pages, provider, language=language, client=client)
